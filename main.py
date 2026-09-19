import os
import re
import fitz
import faiss
import numpy as np

from pathlib import Path
from rank_bm25 import BM25Okapi
from dotenv import load_dotenv
from docx import Document
from google import genai

from langchain_text_splitters import RecursiveCharacterTextSplitter
from sentence_transformers import SentenceTransformer, CrossEncoder


#ENVIRONMENT

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")


#DOCUMENT LOADING

#DOCUMENT LOADING

def load_pdf(file_path):

    document = []

    pdf = fitz.open(file_path)

    for page_number, page in enumerate(
        pdf,
        start=1
    ):

        text = page.get_text().strip()

        if text:

            document.append({
                "text" : text,
                "source" : Path(file_path).name,
                "location" : f"page {page_number}"
            })

    pdf.close()

    return document


def load_docx(file_path):

    document = []

    doc = Document(
        file_path
    )


    #PARAGRAPHS

    for paragraph_number, paragraph in enumerate(
        doc.paragraphs,
        start=1
    ):

        text = paragraph.text.strip()

        if text:

            document.append({
                "text" : text,
                "source" : Path(file_path).name,
                "location" : f"paragraph {paragraph_number}"
            })


    #TABLES

    for table_number, table in enumerate(
        doc.tables,
        start=1
    ):

        rows = []

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:

                    cells.append(
                        text
                    )

            if cells:

                rows.append(
                    " | ".join(cells)
                )


        if rows:

            table_text = "\n".join(
                rows
            )

            document.append({
                "text" : table_text,
                "source" : Path(file_path).name,
                "location" : f"table {table_number}"
            })


    return document


def load_text(file_path):

    document = []

    with open(
        file_path,
        "r",
        encoding="utf-8"
    ) as file:

        text = file.read().strip()


    if text:

        document.append({
            "text" : text,
            "source" : Path(file_path).name,
            "location" : "document"
        })


    return document


def load_document(file_path):

    extension = Path(
        file_path
    ).suffix.lower()


    if extension == ".pdf":

        return load_pdf(
            file_path
        )


    elif extension == ".docx":

        return load_docx(
            file_path
        )


    elif extension in [
        ".txt",
        ".md"
    ]:

        return load_text(
            file_path
        )


    else:

        return []


#CHUNK FUNCTION

#CHUNK FUNCTION

def chunking(documents):

    splitter = RecursiveCharacterTextSplitter(
        chunk_size = 500,
        chunk_overlap = 50
    )

    chunks = []

    chunk_id = 0


    for document in documents:

        chunk_texts = splitter.split_text(
            document["text"]
        )


        for text in chunk_texts:

            chunks.append({
                "text" : text,
                "source" : document["source"],
                "location" : document["location"],
                "chunk_id" : chunk_id
            })

            chunk_id += 1


    return chunks


#MODELS

def load_models():

    model = SentenceTransformer(
        "all-MiniLM-L6-v2"
    )

    reranker = CrossEncoder(
        "cross-encoder/ms-marco-MiniLM-L-6-v2"
    )

    return model, reranker


#BUILD RAG

def build_rag(file_path):

    documents = load_document(
        file_path
    )

    if not documents:

        raise ValueError(
            "No text could be extracted from the document."
        )

    chunks = chunking(
        documents
    )

    if not chunks:

        raise ValueError(
            "No chunks were created from the document."
        )

    model, reranker = load_models()


    #EMBEDDING

    texts = [
        chunk["text"]
        for chunk in chunks
    ]

    embeddings = model.encode(
        texts
    )

    embedding_matrix = np.array(
        embeddings
    ).astype("float32")


    #FAISS

    dimension = embedding_matrix.shape[1]

    index = faiss.IndexFlatL2(
        dimension
    )

    index.add(
        embedding_matrix
    )


    #BM25

    tokenized_texts = [
        text.lower().split()
        for text in texts
    ]

    bm25 = BM25Okapi(
        tokenized_texts
    )


    return {
        "documents" : documents,
        "chunks" : chunks,
        "model" : model,
        "reranker" : reranker,
        "index" : index,
        "bm25" : bm25
    }


#EVIDENCE

def extract_evidence(
    query,
    chunks
):

    query_words = set(
        re.findall(
            r"\b\w+\b",
            query.lower()
        )
    )

    evidence = []

    for chunk in chunks:

        sentences = re.split(
            r"(?<=[.!?])\s+|\n+",
            chunk["text"]
        )

        for sentence in sentences:

            sentence = sentence.strip()

            if not sentence:

                continue

            sentence_words = set(
                re.findall(
                    r"\b\w+\b",
                    sentence.lower()
                )
            )

            overlap = len(
                query_words &
                sentence_words
            )

            evidence.append(
                (
                    overlap,
                    sentence,
                    chunk
                )
            )

    evidence.sort(
        key=lambda x : x[0],
        reverse=True
    )

    return evidence[:3]


#ASK QUESTION

def ask_question(
    rag,
    query
):

    chunks = rag["chunks"]

    model = rag["model"]

    reranker = rag["reranker"]

    index = rag["index"]

    bm25 = rag["bm25"]


    #QUERY EMBEDDING

    query_embedding = model.encode(
        [query]
    )

    query_embedding = np.array(
        query_embedding
    ).astype("float32")


    #FAISS SEARCH

    k = min(
        5,
        len(chunks)
    )

    distance, indices = index.search(
        query_embedding,
        k
    )

    faiss_indices = indices[0].tolist()


    #BM25 SEARCH

    query_tokens = query.lower().split()

    bm25_scores = bm25.get_scores(
        query_tokens
    )

    bm25_indices = sorted(
        range(len(bm25_scores)),
        key=lambda i : bm25_scores[i],
        reverse=True
    )[:k]


    #HYBRID SEARCH

    hybrid_indices = list(
        dict.fromkeys(
            faiss_indices +
            bm25_indices
        )
    )


    retrieved_chunks = [
        chunks[i]
        for i in hybrid_indices
    ]


    #CROSS ENCODER

    pairs = [
        [
            query,
            chunk["text"]
        ]
        for chunk in retrieved_chunks
    ]

    scores = reranker.predict(
        pairs
    )


    ranked_results = sorted(
        zip(
            scores,
            retrieved_chunks
        ),
        key=lambda x : x[0],
        reverse=True
    )


    top_chunks = [
        chunk
        for score, chunk
        in ranked_results[:3]
    ]


    #EVIDENCE

    evidence_results = extract_evidence(
        query,
        top_chunks
    )


    #CONTEXT

    context = "\n\n".join(
        chunk["text"]
        for chunk in top_chunks
    )


    #LLM

    client = genai.Client(
        api_key=GEMINI_API_KEY
    )


    prompt = f"""
Answer the question using only the context below.

Context:
{context}

Question:
{query}

Rules:
- Use only the provided context.
- Do not invent information.
- If the answer is not present, say:
  "The information is not available in the provided document."
- Keep the answer concise and clear.
"""


    try:     

        response = client.models.generate_content(
            model="gemini-3.6-flash",
            contents=prompt
        )

    except Exception as error:
        error_message = str(error).lower()

        if (

            "429" in error_message
            or "quota" in error_message
            or "rate limit" in error_message
            or "resource exhausted" in error_message
        ):
            

            return {
                "answer" : (
                "Gemini API limit has been reached. "
                "Please try again later or use another API key."

                ),

                "citations" : [],
                "chunks" : top_chunks,
                "api_limit" : True
            }

        raise


    #CITATIONS

    citations = []

    for score, sentence, chunk in evidence_results:

        if score > 0:

            citation = {
                "source" : chunk["source"],
                "location" : chunk["location"],
                "text" : sentence
            }

            if citation not in citations:

                citations.append(
                    citation
                )


    return {
        "answer" : response.text,
        "citations" : citations,
        "chunks" : top_chunks
    }